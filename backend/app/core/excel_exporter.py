import io
from typing import List, Dict, Any, Optional

# Lazy import of openpyxl moved inside functions to avoid import errors during test collection

def create_styled_excel(
    title: str,
    subtitle: str,
    columns: List[Dict[str, Any]],  # [{"key": "emp_code", "title": "Mã NV", "width": 12, "format": "text"}, ...]
    data: List[Dict[str, Any]],
    include_totals: bool = False,
    total_keys: Optional[List[str]] = None,
    freeze_cell: Optional[str] = None
) -> io.BytesIO:
    """
    Generate professional styled Excel file with TechZone branding.
    Includes freeze panes, auto-fit column widths, zebra striping, and safe padding.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Báo cáo"
    ws.views.sheetView[0].showGridLines = True

    # Style constants
    HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy Blue
    HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="475569")
    BOLD_FONT = Font(name="Segoe UI", size=10, bold=True)
    REGULAR_FONT = Font(name="Segoe UI", size=10)
    TOTAL_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    THIN_BORDER = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    DOUBLE_BOTTOM_BORDER = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='double', color='0F172A')
    )

    num_cols = len(columns)
    last_col_letter = get_column_letter(max(num_cols, 1))

    # 1. Company Name Header
    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = "TECHZONE - CÔNG TY TNHH TM DV TECH ZONE"
    ws["A1"].font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 20

    # 2. Report Title
    ws.merge_cells(f"A2:{last_col_letter}2")
    ws["A2"] = title.upper()
    ws["A2"].font = TITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 26

    # 3. Subtitle / Timestamp
    ws.merge_cells(f"A3:{last_col_letter}3")
    ws["A3"] = subtitle
    ws["A3"].font = SUBTITLE_FONT
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 18

    # Blank buffer row
    ws.row_dimensions[4].height = 10

    row_idx = 5

    # 4. Table Headers (Height 36 so 2-line titles are never cut off)
    ws.row_dimensions[row_idx].height = 36
    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.value = col.get("title", col.get("key"))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    row_idx += 1

    # 5. Data Rows with Alternating Zebra Striping
    for idx, item in enumerate(data):
        ws.row_dimensions[row_idx].height = 23
        row_fill = ZEBRA_FILL if idx % 2 == 1 else WHITE_FILL

        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = item.get(col["key"])
            fmt = col.get("format", "text")

            if val is None:
                cell.value = ""
            elif fmt == "currency":
                cell.value = float(val)
                cell.number_format = '#,##0 "đ"'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif fmt == "number":
                cell.value = float(val) if isinstance(val, (int, float)) else val
                cell.number_format = '#,##0.0' if isinstance(cell.value, float) and not cell.value.is_integer() else '#,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif fmt == "percentage":
                cell.value = float(val) / 100.0 if float(val) > 1 else float(val)
                cell.number_format = '0.0%'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.value = str(val)
                cell.alignment = Alignment(horizontal="center" if col.get("align") == "center" else "left", vertical="center")

            cell.font = REGULAR_FONT
            cell.fill = row_fill
            cell.border = THIN_BORDER
        row_idx += 1

    # 6. Totals Row
    if include_totals and total_keys:
        ws.row_dimensions[row_idx].height = 26
        first_total_col = True
        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = TOTAL_FILL
            cell.font = BOLD_FONT
            cell.border = DOUBLE_BOTTOM_BORDER

            if col["key"] in total_keys:
                total_val = sum(float(item.get(col["key"]) or 0) for item in data)
                cell.value = total_val
                if col.get("format") == "currency":
                    cell.number_format = '#,##0 "đ"'
                elif col.get("format") == "number":
                    cell.number_format = '#,##0.0' if not float(total_val).is_integer() else '#,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif first_total_col:
                cell.value = "TỔNG CỘNG"
                cell.alignment = Alignment(horizontal="left", vertical="center")
                first_total_col = False

    # 7. Freeze Panes (Row 5 header + first frozen columns so scrolling never loses context)
    if freeze_cell:
        ws.freeze_panes = freeze_cell
    elif num_cols >= 4:
        if columns and columns[0].get("key") in ["salary_period_label", "period", "salary_period"]:
            ws.freeze_panes = "B6"
        else:
            ws.freeze_panes = "C6"
    else:
        ws.freeze_panes = "A6"

    # 8. Set Dynamic Safe Column Widths
    for col_idx, col in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        header_title = str(col.get("title", col.get("key", "")))

        # Longest word in header (to prevent word wrap overflow)
        words = header_title.replace("/", " ").replace("(", " ").replace(")", " ").split()
        max_word_len = max([len(w) for w in words] + [0])

        # Check data length in this column
        data_lens = []
        for r in range(6, row_idx):
            v = ws.cell(row=r, column=col_idx).value
            if v is not None:
                data_lens.append(len(str(v)))
        max_data_len = max(data_lens + [0])

        # Currency format adds ~4-5 characters (commas, spaces, đ)
        if col.get("format") == "currency":
            max_data_len += 5

        base_width = col.get("width", 14)
        # Generous margin so headers and currency numbers never show ### or get clipped
        optimal_width = max(base_width, max_word_len + 4, max_data_len + 3, len(header_title) if len(header_title) < 18 else 14)
        ws.column_dimensions[col_letter].width = min(max(optimal_width, 13), 48)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def create_multisheet_excel(sheets: List[Dict[str, Any]]) -> io.BytesIO:
    """
    Generate professional multi-sheet Excel file with TechZone branding.
    Includes freeze panes, auto-fit column widths, zebra striping, and safe padding.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    default_sheet = wb.active

    HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy Blue
    HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="475569")
    BOLD_FONT = Font(name="Segoe UI", size=10, bold=True)
    REGULAR_FONT = Font(name="Segoe UI", size=10)
    TOTAL_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    THIN_BORDER = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    DOUBLE_BOTTOM_BORDER = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='double', color='0F172A')
    )

    for idx, s in enumerate(sheets):
        sheet_title = s.get("sheet_name", f"Sheet {idx + 1}")[:30]
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        title = s.get("title", "")
        subtitle = s.get("subtitle", "")
        columns = s.get("columns", [])
        data = s.get("data", [])
        include_totals = s.get("include_totals", False)
        total_keys = s.get("total_keys", [])

        num_cols = len(columns)
        last_col_letter = get_column_letter(max(num_cols, 1))

        # 1. Company Name Header
        ws.merge_cells(f"A1:{last_col_letter}1")
        ws["A1"] = "TECHZONE - CÔNG TY TNHH TM DV TECH ZONE"
        ws["A1"].font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 20

        # 2. Report Title
        ws.merge_cells(f"A2:{last_col_letter}2")
        ws["A2"] = title.upper()
        ws["A2"].font = TITLE_FONT
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 26

        # 3. Subtitle / Timestamp
        ws.merge_cells(f"A3:{last_col_letter}3")
        ws["A3"] = subtitle
        ws["A3"].font = SUBTITLE_FONT
        ws["A3"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[3].height = 18

        ws.row_dimensions[4].height = 10
        row_idx = 5

        # 4. Table Headers
        ws.row_dimensions[row_idx].height = 36
        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = col.get("title", col.get("key"))
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN_BORDER

        row_idx += 1

        # 5. Data Rows
        for d_idx, item in enumerate(data):
            ws.row_dimensions[row_idx].height = 23
            row_fill = ZEBRA_FILL if d_idx % 2 == 1 else WHITE_FILL

            for col_idx, col in enumerate(columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                val = item.get(col["key"])
                fmt = col.get("format", "text")

                if val is None:
                    cell.value = ""
                elif fmt == "currency":
                    cell.value = float(val)
                    cell.number_format = '#,##0 "đ"'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif fmt == "number":
                    cell.value = float(val) if isinstance(val, (int, float)) else val
                    cell.number_format = '#,##0.0' if isinstance(cell.value, float) and not cell.value.is_integer() else '#,##0'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif fmt == "percentage":
                    cell.value = float(val) / 100.0 if float(val) > 1 else float(val)
                    cell.number_format = '0.0%'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.value = str(val)
                    cell.alignment = Alignment(horizontal="center" if col.get("align") == "center" else "left", vertical="center")

                cell.font = REGULAR_FONT
                cell.fill = row_fill
                cell.border = THIN_BORDER
            row_idx += 1

        # 6. Totals Row
        if include_totals and total_keys:
            ws.row_dimensions[row_idx].height = 26
            first_total_col = True
            for col_idx, col in enumerate(columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = TOTAL_FILL
                cell.font = BOLD_FONT
                cell.border = DOUBLE_BOTTOM_BORDER

                if col["key"] in total_keys:
                    total_val = sum(float(item.get(col["key"]) or 0) for item in data)
                    cell.value = total_val
                    if col.get("format") == "currency":
                        cell.number_format = '#,##0 "đ"'
                    elif col.get("format") == "number":
                        cell.number_format = '#,##0.0' if not float(total_val).is_integer() else '#,##0'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif first_total_col:
                    cell.value = "TỔNG CỘNG"
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    first_total_col = False

        # 7. Freeze Panes
        freeze_cell = s.get("freeze_cell")
        if freeze_cell:
            ws.freeze_panes = freeze_cell
        elif num_cols >= 4:
            if columns and columns[0].get("key") in ["salary_period_label", "period", "salary_period"]:
                ws.freeze_panes = "B6"
            else:
                ws.freeze_panes = "C6"
        else:
            ws.freeze_panes = "A6"

        # 8. Set Column Widths
        for col_idx, col in enumerate(columns, start=1):
            col_letter = get_column_letter(col_idx)
            header_title = str(col.get("title", col.get("key", "")))

            words = header_title.replace("/", " ").replace("(", " ").replace(")", " ").split()
            max_word_len = max([len(w) for w in words] + [0])

            data_lens = []
            for r in range(6, row_idx):
                v = ws.cell(row=r, column=col_idx).value
                if v is not None:
                    data_lens.append(len(str(v)))
            max_data_len = max(data_lens + [0])

            if col.get("format") == "currency":
                max_data_len += 5

            base_width = col.get("width", 14)
            optimal_width = max(base_width, max_word_len + 4, max_data_len + 3, len(header_title) if len(header_title) < 18 else 14)
            ws.column_dimensions[col_letter].width = min(max(optimal_width, 13), 48)

    # Remove the initial blank sheet
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
